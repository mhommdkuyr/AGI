#!/usr/bin/env python3
"""Audit an Ibb business-directory XLSX without treating its fields as verified.

The tool is intentionally offline and read-only with respect to the input workbook.
It produces aggregate quality metrics and a CSV of rows that need review. It does
not verify businesses against Google Maps, infer missing facts, or remove source rows.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - environment-specific
    raise SystemExit("Missing dependency: openpyxl. Install with: python -m pip install openpyxl") from exc


ALIASES = {
    "name": ("اسم المنشأة أو المعلم", "اسم المنشأة", "اسم النشاط", "Business Name", "Name"),
    "category": ("التصنيف", "الفئة", "النشاط", "Category"),
    "address": ("العنوان التفصيلي والشارع", "العنوان", "Address"),
    "phone": ("الهاتف", "رقم الهاتف", "Phone", "Telephone"),
    "lat": ("خط العرض Lat", "خط العرض", "Latitude", "lat"),
    "lng": ("خط الطول Lng", "خط الطول", "Longitude", "lng", "lon"),
    "maps": ("رابط الموقع على خريطة جوجل", "رابط الخريطة", "رابط الموقع", "Google Maps URL", "Maps URL"),
    "rating": ("التقييم", "Rating"),
    "floors": ("الطوابق", "Floors"),
    "height": ("الارتفاع م", "الارتفاع", "Height"),
}
TEMPLATE_PATTERNS = (
    re.compile(r"#\s*\d+\s*$"),
    re.compile(r"رقم\s*\d+\s*$"),
)


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalize(value: Any) -> str:
    value = unicodedata.normalize("NFKC", text(value)).lower()
    value = re.sub(r"[\u064b-\u065f\u0670]", "", value)
    value = re.sub(r"[إأآٱ]", "ا", value).replace("ى", "ي").replace("ة", "ه")
    value = re.sub(r"[^\w\u0600-\u06ff]+", " ", value, flags=re.UNICODE)
    return re.sub(r"\s+", " ", value).strip()


def number(value: Any) -> float | None:
    try:
        result = float(text(value).replace(",", "."))
        return result if result == result else None
    except (TypeError, ValueError):
        return None


def pick_headers(headers: list[str]) -> dict[str, int | None]:
    normalized = {normalize(header): index for index, header in enumerate(headers)}
    chosen: dict[str, int | None] = {}
    for field, aliases in ALIASES.items():
        idx = None
        for alias in aliases:
            key = normalize(alias)
            if key in normalized:
                idx = normalized[key]
                break
        chosen[field] = idx
    return chosen


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Input XLSX file")
    parser.add_argument("--sheet", default=None, help="Worksheet name; default is the first sheet")
    parser.add_argument("--report", type=Path, default=Path("ibb_inventory_audit.md"))
    parser.add_argument("--flagged-csv", type=Path, default=Path("ibb_inventory_flagged_rows.csv"))
    args = parser.parse_args()

    if not args.workbook.is_file():
        parser.error(f"Workbook not found: {args.workbook}")

    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    if args.sheet:
        if args.sheet not in wb.sheetnames:
            parser.error(f"Sheet not found: {args.sheet}. Available sheets: {', '.join(wb.sheetnames)}")
        ws = wb[args.sheet]
    else:
        ws = wb[wb.sheetnames[0]]

    iterator = ws.iter_rows(values_only=True)
    try:
        first_row = next(iterator)
    except StopIteration:
        parser.error("Worksheet is empty.")

    headers = [text(value) for value in first_row]
    columns = pick_headers(headers)
    if columns["name"] is None:
        parser.error("Could not identify a business-name column. Expected one of: " + ", ".join(ALIASES["name"]))

    records: list[dict[str, Any]] = []
    exact_rows = Counter()
    normalized_names = Counter()
    categories = Counter()
    generic_map_count = 0
    template_name_count = 0
    template_address_count = 0
    phone_7plus_count = 0
    rating_count = floors_count = height_count = 0
    coordinate_counts: Counter[tuple[float, float]] = Counter()
    invalid_coordinate_rows = 0

    for excel_row, values in enumerate(iterator, start=2):
        row = list(values) + [None] * max(0, len(headers) - len(values))
        record = {header: row[idx] if idx < len(row) else None for idx, header in enumerate(headers)}
        records.append(record)
        exact_rows[tuple(text(value) for value in row[:len(headers)])] += 1

        name = text(row[columns["name"]])
        address = text(row[columns["address"]]) if columns["address"] is not None else ""
        phone = text(row[columns["phone"]]) if columns["phone"] is not None else ""
        category = text(row[columns["category"]]) if columns["category"] is not None else ""
        map_url = text(row[columns["maps"]]) if columns["maps"] is not None else ""

        if name:
            normalized_names[normalize(name)] += 1
        if category:
            categories[category] += 1
        if any(pattern.search(name) for pattern in TEMPLATE_PATTERNS):
            template_name_count += 1
        if re.search(r"مبنى\s*رقم", address) and "الشارع الفرعي" in address:
            template_address_count += 1
        if re.search(r"maps/search/\?api=1&query=", map_url, flags=re.IGNORECASE):
            generic_map_count += 1
        digits = re.sub(r"\D", "", phone)
        if len(digits) >= 7:
            phone_7plus_count += 1
        if columns["rating"] is not None and number(row[columns["rating"]]) is not None:
            rating_count += 1
        if columns["floors"] is not None and text(row[columns["floors"]]):
            floors_count += 1
        if columns["height"] is not None and text(row[columns["height"]]):
            height_count += 1

        lat = number(row[columns["lat"]]) if columns["lat"] is not None else None
        lng = number(row[columns["lng"]]) if columns["lng"] is not None else None
        if lat is not None and lng is not None:
            coordinate_counts[(lat, lng)] += 1
        elif columns["lat"] is not None or columns["lng"] is not None:
            invalid_coordinate_rows += 1

    duplicate_rows = sum(count - 1 for count in exact_rows.values() if count > 1)
    duplicate_coordinate_rows = sum(count for count in coordinate_counts.values() if count > 1)
    repeated_names = sum(count for count in normalized_names.values() if count > 1)
    data_rows = len(records)
    collected_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    flags: list[dict[str, Any]] = []
    for index, record in enumerate(records, start=2):
        row_values = [record.get(header) for header in headers]
        name = text(record.get(headers[columns["name"]]))
        address = text(record.get(headers[columns["address"]])) if columns["address"] is not None else ""
        map_url = text(record.get(headers[columns["maps"]])) if columns["maps"] is not None else ""
        reasons = []
        if any(pattern.search(name) for pattern in TEMPLATE_PATTERNS):
            reasons.append("name_has_numbered_template_suffix")
        if re.search(r"مبنى\s*رقم", address) and "الشارع الفرعي" in address:
            reasons.append("address_has_template_pattern")
        if map_url and re.search(r"maps/search/\?api=1&query=", map_url, flags=re.IGNORECASE):
            reasons.append("generic_coordinate_search_url_not_place_permalink")
        if reasons:
            flags.append({
                "excel_row": index,
                "business_name": name,
                "review_flags": "|".join(reasons),
                "source_map_url": map_url,
                "source_address": address,
            })

    args.flagged_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.flagged_csv.open("w", encoding="utf-8-sig", newline="") as output:
        fields = ["excel_row", "business_name", "review_flags", "source_map_url", "source_address"]
        writer = csv.DictWriter(output, fieldnames=fields)
        writer.writeheader()
        writer.writerows(flags)

    summary = {
        "workbook": str(args.workbook),
        "worksheet": ws.title,
        "checked_at": collected_at,
        "rows": data_rows,
        "columns": len(headers),
        "exact_duplicate_rows_after_first_occurrence": duplicate_rows,
        "unique_normalized_names": len(normalized_names),
        "repeated_name_rows": repeated_names,
        "categories": len(categories),
        "numbered_template_names": template_name_count,
        "template_address_patterns": template_address_count,
        "generic_google_maps_coordinate_search_urls": generic_map_count,
        "phone_fields_with_7plus_digits_after_separator_removal": phone_7plus_count,
        "numeric_ratings": rating_count if columns["rating"] is not None else None,
        "nonempty_floor_fields": floors_count if columns["floors"] is not None else None,
        "nonempty_height_fields": height_count if columns["height"] is not None else None,
        "unique_coordinate_pairs": len(coordinate_counts),
        "rows_in_repeated_coordinate_groups": duplicate_coordinate_rows,
        "rows_with_missing_or_non_numeric_coordinates": invalid_coordinate_rows,
        "flagged_rows": len(flags),
        "flagged_csv": str(args.flagged_csv),
        "warning": "Automated flags are review leads, not proof of fabricated data. A filled field is not the same as a verified field.",
    }

    report = [
        "# Ibb Inventory Quality Audit",
        "",
        f"- Workbook: {args.workbook}",
        f"- Worksheet: {ws.title}",
        f"- Checked at: {collected_at}",
        f"- Data rows: {data_rows:,}",
        f"- Columns: {len(headers)}",
        f"- Exact duplicate rows after first occurrence: {duplicate_rows:,}",
        f"- Unique normalized names: {len(normalized_names):,}",
        f"- Repeated-name rows: {repeated_names:,}",
        f"- Categories: {len(categories):,}",
        f"- Numbered-template names: {template_name_count:,}",
        f"- Template-like address patterns: {template_address_count:,}",
        f"- Generic Google Maps coordinate-search URLs: {generic_map_count:,}",
        f"- Phone fields with 7+ digits after separator removal: {phone_7plus_count:,}",
        f"- Rows in repeated coordinate groups: {duplicate_coordinate_rows:,}",
        f"- Missing/non-numeric coordinate rows: {invalid_coordinate_rows:,}",
        f"- Flagged rows for review: {len(flags):,}",
        "",
        "## Important interpretation",
        "",
        "This report is a quality screen, not business verification. Template patterns are heuristics, phone-like digits do not prove that a number belongs to the business, and a generic coordinate search URL does not establish place identity. Preserve the input data; check each listing against a source page and a stable location permalink before publication.",
        "",
        "## Category counts (top 30)",
        "",
    ]
    for category, count in categories.most_common(30):
        report.append(f"- {category or '(blank)'}: {count:,}")

    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text("\n".join(report) + "\n", encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
